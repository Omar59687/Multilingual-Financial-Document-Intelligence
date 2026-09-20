"""Native parsers: one function per format, same Document contract.

No OCR, no guessing: tables only via pdfplumber lattice (ruled grids);
numbers stay native; text preserved verbatim (stripped of surrounding
whitespace only).
"""

import csv
from pathlib import Path

from .model import (Document, Element, ElementType, FileType,
                    IngestionStatus, LanguageHint)

# Transparent, documented thresholds (see docs/INGESTION_DESIGN.md).
NATIVE_PAGE_MIN_CHARS = 50   # pdf page counts as natively sufficient at/above
MIXED_MIN_SHARE = 0.20       # script share needed for ar/en/mixed hints


def language_hint(text: str) -> LanguageHint:
    """Unicode-script-ratio hint (heuristic, NOT language detection)."""
    if not text or not text.strip():
        return LanguageHint.UNKNOWN
    arabic = sum(1 for ch in text if "\u0600" <= ch <= "\u06FF"
                 or "\u0750" <= ch <= "\u077F" or "\uFB50" <= ch <= "\uFDFF"
                 or "\uFE70" <= ch <= "\uFEFF")
    latin = sum(1 for ch in text if "A" <= ch <= "Z" or "a" <= ch <= "z")
    total = max(len(text), 1)
    has_ar, has_en = arabic / total >= MIXED_MIN_SHARE, \
        latin / total >= MIXED_MIN_SHARE
    if has_ar and has_en:
        return LanguageHint.MIXED
    if has_ar:
        return LanguageHint.AR
    if has_en:
        return LanguageHint.EN
    return LanguageHint.UNKNOWN


def _lib_version(module_name: str) -> str:
    try:
        module = __import__(module_name)
        return getattr(module, "__version__", "?")
    except Exception:
        return "missing"


# ---------------------------------------------------------------------------
# PDF (pypdf text + presence; pdfplumber lattice tables only)
# ---------------------------------------------------------------------------
def _page_has_images(page) -> bool:
    try:
        resources = page.get("/Resources")
        xobjects = resources.get("/XObject") if resources else None
    except Exception:
        return False
    if not xobjects:
        return False
    try:
        for obj in xobjects.values():
            try:
                if obj.get_object().get("/Subtype") == "/Image":
                    return True
            except Exception:
                continue
    except Exception:
        return False
    return False


def _lattice_tables(path: Path, page_number: int):
    """Ruled-grid tables via pdfplumber lattice strategy. Returns a list of
    row-lists (strings). Empty list = none reliably found (deferred)."""
    import pdfplumber
    with pdfplumber.open(str(path)) as pdf:
        if page_number - 1 >= len(pdf.pages):
            return []
        page = pdf.pages[page_number - 1]
        found = page.find_tables(
            table_settings={"vertical_strategy": "lines",
                            "horizontal_strategy": "lines",
                            "snap_tolerance": 3})
        # pdfplumber >= 0.11 unified API: Table.extract() (list of rows).
        return [table.extract() or [] for table in found]


def parse_pdf(path: Path, document_id: str) -> Document:
    from pypdf import PdfReader
    path = Path(path)
    doc = Document(document_id=document_id, filename=path.name,
                   file_type=FileType.PDF,
                   metadata={"parser": "pypdf+" + _lib_version("pypdf")
                             + "/pdfplumber-" + _lib_version("pdfplumber"),
                             "table_strategy": "lattice-only"})
    reader = PdfReader(str(path))
    doc.page_count = len(reader.pages)
    pages_meta, full_text, table_total = [], [], 0
    text_counter = 0
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            text = ""
            pages_meta.append({"page": i, "chars": 0, "has_images": False,
                               "native_sufficient": False,
                               "extract_error": str(exc)})
            continue
        stripped = text.strip()
        images = _page_has_images(page)
        sufficient = len(stripped) >= NATIVE_PAGE_MIN_CHARS
        pages_meta.append({"page": i, "chars": len(stripped),
                           "has_images": images,
                           "native_sufficient": sufficient})
        if stripped:
            text_counter += 1
            doc.elements.append(Element(
                element_id=f"{document_id}:p{i}:text:{text_counter}",
                element_type=ElementType.TEXT, page=i, text=stripped,
                meta={"chars": len(stripped)}))
            full_text.append(stripped)
        try:
            tables = _lattice_tables(path, i)
        except Exception as exc:
            tables = []
            pages_meta[-1]["table_error"] = str(exc)
        for ti, rows in enumerate(tables):
            table_total += 1
            doc.elements.append(Element(
                element_id=f"{document_id}:p{i}:table:{ti}",
                element_type=ElementType.TABLE, page=i, table_index=ti,
                meta={"n_rows": len(rows),
                      "n_cols": max((len(r) for r in rows), default=0)}))
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    doc.elements.append(Element(
                        element_id=f"{document_id}:p{i}:cell:{ti}r{r}c{c}",
                        element_type=ElementType.TABLE_CELL, page=i,
                        row=r, column=c, table_index=ti,
                        text=(cell or "").strip()))
    doc.metadata["pages"] = pages_meta
    doc.metadata["tables_found"] = table_total
    doc.language_hint = language_hint("\n".join(full_text))
    sufficient_pages = sum(1 for p in pages_meta if p["native_sufficient"])
    if sufficient_pages == len(pages_meta) and pages_meta:
        doc.status = IngestionStatus.NATIVE_OK
    elif sufficient_pages == 0:
        if any(p["has_images"] for p in pages_meta):
            doc.status = IngestionStatus.REQUIRES_OCR
            doc.requires_ocr = True
        else:
            doc.status = IngestionStatus.FAILED
            doc.error = "no native text and no page images"
    else:
        doc.status = IngestionStatus.PARTIAL_NATIVE
        doc.requires_ocr = True
    return doc


# ---------------------------------------------------------------------------
# DOCX (python-docx)
# ---------------------------------------------------------------------------
def parse_docx(path: Path, document_id: str) -> Document:
    from docx import Document as DocxDocument
    path = Path(path)
    doc = Document(document_id=document_id, filename=path.name,
                   file_type=FileType.DOCX,
                   metadata={"parser": "python-docx-"
                             + _lib_version("docx")})
    ox = DocxDocument(str(path))
    texts, skipped_empty, para_n = [], 0, 0
    for para in ox.paragraphs:
        text = para.text.strip()
        if not text:
            skipped_empty += 1
            continue
        para_n += 1
        is_heading = (para.style.name or "").startswith("Heading")
        texts.append(text)
        doc.elements.append(Element(
            element_id=f"{document_id}:para:{para_n}",
            element_type=ElementType.HEADING if is_heading
            else ElementType.TEXT,
            text=text,
            meta={"style": para.style.name, "order": para_n,
                  "rtl": bool(para._p.xpath(".//w:bidi"))}))
    for ti, table in enumerate(ox.tables):
        doc.elements.append(Element(
            element_id=f"{document_id}:table:{ti}",
            element_type=ElementType.TABLE, table_index=ti,
            meta={"n_rows": len(table.rows),
                  "n_cols": max((len(r.cells) for r in table.rows),
                                default=0)}))
        for r, row in enumerate(table.rows):
            doc.elements.append(Element(
                element_id=f"{document_id}:t{ti}r{r}",
                element_type=ElementType.TABLE_ROW, table_index=ti, row=r))
            for c, cell in enumerate(row.cells):
                doc.elements.append(Element(
                    element_id=f"{document_id}:t{ti}r{r}c{c}",
                    element_type=ElementType.TABLE_CELL, table_index=ti,
                    row=r, column=c, text=cell.text.strip()))
    for si, _shape in enumerate(ox.inline_shapes):
        doc.elements.append(Element(
            element_id=f"{document_id}:img:{si}",
            element_type=ElementType.IMAGE_REF,
            meta={"scope": "document", "index": si,
                  "note": "embedded image not processed in Phase 1"}))
    doc.metadata.update({"paragraphs": para_n,
                         "skipped_empty_paragraphs": skipped_empty,
                         "tables": len(ox.tables),
                         "inline_images": len(ox.inline_shapes)})
    doc.language_hint = language_hint("\n".join(texts))
    if para_n or ox.tables:
        doc.status = IngestionStatus.NATIVE_OK
    else:
        doc.status = IngestionStatus.FAILED
        doc.error = "no paragraphs or tables found"
    return doc


# ---------------------------------------------------------------------------
# XLSX (openpyxl, read-only; native values preserved)
# ---------------------------------------------------------------------------
def parse_xlsx(path: Path, document_id: str) -> Document:
    from openpyxl import load_workbook
    path = Path(path)
    doc = Document(document_id=document_id, filename=path.name,
                   file_type=FileType.XLSX,
                   metadata={"parser": "openpyxl-"
                             + _lib_version("openpyxl")})
    wb = load_workbook(str(path), read_only=True, data_only=False)
    sheets_meta, texts, data_rows = [], [], 0
    for ws in wb.worksheets:
        try:
            merged = [str(r) for r in ws.merged_cells.ranges]
        except AttributeError:
            merged = []  # read-only worksheets expose no merged_cells
        sheets_meta.append({"name": ws.title, "max_row": ws.max_row,
                            "max_column": ws.max_column, "merged": merged})
        for row in ws.iter_rows(values_only=False):
            for cell in row:
                if cell.value is None:
                    continue
                data_rows += 1
                if cell.data_type == "f":
                    value, meta = None, {"formula": cell.value,
                                         "note": "no cached value in Phase 1"}
                else:
                    value, meta = cell.value, {}
                    if isinstance(value, str):
                        texts.append(value)
                doc.elements.append(Element(
                    element_id=f"{document_id}:{ws.title}:"
                               f"r{cell.row}c{cell.column}",
                    element_type=ElementType.SHEET_CELL, sheet=ws.title,
                    row=cell.row, column=cell.column,
                    text=str(value) if isinstance(value, str) else None,
                    value=value, data_type=cell.data_type,
                    meta={"coordinate": cell.coordinate, **meta}))
    wb.close()
    doc.sheet_count = len(wb.sheetnames)
    doc.metadata["sheets"] = sheets_meta
    doc.metadata["data_cells"] = data_rows
    doc.language_hint = language_hint("\n".join(texts))
    if data_rows:
        doc.status = IngestionStatus.NATIVE_OK
    else:
        doc.status = IngestionStatus.FAILED
        doc.error = "no data cells found"
    return doc


# ---------------------------------------------------------------------------
# CSV (stdlib; utf-8-sig BOM handling)
# ---------------------------------------------------------------------------
def parse_csv(path: Path, document_id: str) -> Document:
    path = Path(path)
    doc = Document(document_id=document_id, filename=path.name,
                   file_type=FileType.CSV,
                   metadata={"parser": "stdlib-csv", "encoding": "utf-8-sig",
                             "delimiter": ","})
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        doc.metadata["header"] = list(reader.fieldnames or [])
        texts = []
        for n, row in enumerate(reader, start=1):
            mapping = {k: (v if v is not None else "")
                       for k, v in row.items()}
            texts.extend(v for v in mapping.values() if v)
            doc.elements.append(Element(
                element_id=f"{document_id}:row:{n}",
                element_type=ElementType.CSV_ROW, row=n,
                value=mapping, data_type="str",
                meta={"line": n + 1}))  # +1 for header line
    doc.metadata["data_rows"] = len(doc.elements)
    doc.language_hint = language_hint("\n".join(texts))
    if doc.elements:
        doc.status = IngestionStatus.NATIVE_OK
    else:
        doc.status = IngestionStatus.FAILED
        doc.error = "no data rows found"
    return doc


# ---------------------------------------------------------------------------
# Images (Pillow metadata only — content NEVER interpreted here)
# ---------------------------------------------------------------------------
def parse_image(path: Path, document_id: str, file_type) -> Document:
    from PIL import Image
    path = Path(path)
    doc = Document(document_id=document_id, filename=path.name,
                   file_type=file_type,
                   status=IngestionStatus.REQUIRES_VISUAL,
                   requires_visual=True,
                   metadata={"parser": "Pillow-" + _lib_version("PIL")})
    with Image.open(path) as img:
        img.load()
        doc.metadata.update({"format": img.format, "width": img.width,
                             "height": img.height, "mode": img.mode})
    doc.elements.append(Element(
        element_id=f"{document_id}:img:0",
        element_type=ElementType.IMAGE_REF,
        meta={"scope": "whole-file", "note": "visual processing deferred"}))
    return doc
